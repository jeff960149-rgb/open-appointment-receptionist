import csv
import io
from datetime import date, datetime, time

from openpyxl import load_workbook


REQUIRED_COLUMNS = {"date", "therapist", "start_time", "end_time"}
ALLOWED_STATUS = {"available", "open", "yes", "1"}


def _normalize_header(value: object) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def _normalize_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Invalid date: {text!r}. Use YYYY-MM-DD.")


def _normalize_time(value: object) -> str:
    if isinstance(value, datetime):
        return value.time().strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    text = str(value or "").strip()
    for pattern in ("%H:%M", "%I:%M %p", "%I %p"):
        try:
            return datetime.strptime(text.upper(), pattern).time().strftime("%H:%M")
        except ValueError:
            pass
    raise ValueError(f"Invalid time: {text!r}. Use HH:MM or 3:00 PM.")


def _normalize_rows(headers: list[object], rows: list[list[object]]) -> list[dict[str, str]]:
    normalized_headers = [_normalize_header(value) for value in headers]
    missing = REQUIRED_COLUMNS - set(normalized_headers)
    if missing:
        raise ValueError("Missing columns: " + ", ".join(sorted(missing)))

    result: list[dict[str, str]] = []
    for row_number, values in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in values):
            continue
        record = dict(zip(normalized_headers, values))
        therapist = str(record.get("therapist") or "").strip()
        if not therapist:
            raise ValueError(f"Row {row_number}: therapist is required.")
        status = str(record.get("status") or "available").strip().lower()
        start_time = _normalize_time(record.get("start_time"))
        end_time = _normalize_time(record.get("end_time"))
        if end_time <= start_time:
            raise ValueError(f"Row {row_number}: end_time must be after start_time.")
        result.append(
            {
                "date": _normalize_date(record.get("date")),
                "therapist": therapist,
                "start_time": start_time,
                "end_time": end_time,
                "status": "available" if status in ALLOWED_STATUS else "unavailable",
            }
        )
    if not result:
        raise ValueError("The schedule has no data rows.")
    return result


def parse_schedule(filename: str, content: bytes) -> list[dict[str, str]]:
    lower_name = filename.lower()
    if lower_name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = list(csv.reader(io.StringIO(text)))
        if not reader:
            raise ValueError("The CSV file is empty.")
        return _normalize_rows(reader[0], reader[1:])

    if lower_name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        if not values:
            raise ValueError("The Excel file is empty.")
        return _normalize_rows(list(values[0]), [list(row) for row in values[1:]])

    raise ValueError("Only .xlsx and .csv files are supported.")


def normalize_schedule_records(records: list[dict[str, object]]) -> list[dict[str, str]]:
    if not records:
        raise ValueError("The schedule has no data rows.")
    headers = ["date", "therapist", "start_time", "end_time", "status"]
    rows = [[record.get(header) for header in headers] for record in records]
    return _normalize_rows(headers, rows)
